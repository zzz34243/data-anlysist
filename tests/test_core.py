import os
import tempfile
import time
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

from app.db import Database
from app.memory import MemoryStore
from app.orchestrator import Orchestrator
from app.retention import RetentionManager
from app.selection import devote, filter_data, resolve_selection
from app.excel_reader import read_excel_rows


class FixedLLM:
    provider_name = "test-provider"
    model = "test-model"

    def complete(self, *, system, prompt, context):
        return '{"strategy_summary":"先验证重点产品方案，再根据销售结果扩大。","actions":[{"priority":"优先执行","title":"重点产品组合测试","goal":"提高下一周期销售额","why":"产品 A 当前销售额最高。","steps":["设计两种组合方案","比较同周期销售额"],"audience":"目标人群数据不足，需补充客户标签","timing":"下一完整销售周期","channel":"渠道数据不足，需结合实际渠道安排","metrics":["下一周期销售额","产品 A 销售额占比"],"caveat":"先小范围测试"}],"risks":[{"risk":"缺少渠道数据","response":"补充渠道字段后复核"}]}'


class CoreTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp())
        for name in ("runs", "history", "charts", "reports", "logs"): (self.root / name).mkdir()
        self.memory = MemoryStore(Database(self.root / "app.db"), self.root)
        self.orchestrator = Orchestrator(self.memory, self.root, max_retries=1)

    def test_prompt_filter_and_devote(self):
        optimized = self.orchestrator.optimize(request="模糊分析", product_mode="online")
        self.assertIn("线性趋势图", optimized["optimized_prompt"])
        self.assertEqual(len(devote()), 7)
        result = filter_data([{"date":"2026-01-01"}] * 10, selection=resolve_selection(), filters={"limit": 3}, max_rows=5)
        self.assertEqual(len(result["rows"]), 3); self.assertTrue(result["truncated"])

    def test_pipeline_and_trace(self):
        result = self.orchestrator.analyze(project_name="测试", request="分析销售", data=[{"date":"2026/01/01","product":"A","channel":"线上","amount":"100"},{"date":"2026/01/02","product":"A","channel":"线上","amount":"150"}])
        self.assertEqual(result["status"], "completed"); self.assertTrue(result["report_validation"]["valid"]); self.assertEqual(set(result["charts"]["charts"]), {"bar", "line"})
        self.assertEqual([x["step_order"] for x in self.orchestrator.trace(result["run_id"])], list(range(1, 8)))
        self.assertEqual(len(self.memory.run_memory(result["run_id"])["findings"]), 7)
        replay = self.orchestrator.rerun(result["run_id"], from_step=4)
        self.assertEqual(replay["replay"]["source_run_id"], result["run_id"])

    def test_permissions_privacy_and_retention(self):
        project = self.memory.create_project("权限"); run = self.memory.create_run(project, {})
        with self.assertRaises(PermissionError): self.memory.create_plan(actor="insights", project_id=project, run_id=run, objective="x", optimized_prompt="x", steps=[], selection={})
        with self.assertRaises(PermissionError): self.orchestrator.analyze(project_name="隐私", request="x", data=[{"email":"x","amount":1}], privacy_consent=False)
        old = self.root / "runs" / "old.json"; old.write_text("{}"); stamp = time.time() - 8 * 86400; os.utime(old, (stamp, stamp)); self.assertIn(str(old), RetentionManager(self.root, 7).cleanup())

    def test_excel_reader_uses_first_non_empty_sheet(self):
        workbook = Workbook()
        workbook.active.title = "说明"
        sheet = workbook.create_sheet("销售数据")
        sheet.append(["日期", "产品", "渠道", "销量", "单价"])
        sheet.append(["2026-01-01", "A", "线上", 10, 120])
        sheet.append(["2026-02-01", "A", "门店", 8, 150])
        content = BytesIO(); workbook.save(content); workbook.close()
        rows, sheet_name = read_excel_rows(content.getvalue(), "sales.xlsx")
        self.assertEqual(sheet_name, "销售数据")
        self.assertEqual(len(rows), 2)
        result = self.orchestrator.analyze(project_name="Excel 测试", request="分析销售", data=rows)
        self.assertEqual(result["prepared_summary"]["total_amount"], 2400)

    def test_excel_reader_rejects_missing_sales_columns(self):
        workbook = Workbook(); sheet = workbook.active
        sheet.append(["产品", "备注"]); sheet.append(["A", "无日期和金额"])
        content = BytesIO(); workbook.save(content); workbook.close()
        with self.assertRaisesRegex(ValueError, "缺少必要列"):
            read_excel_rows(content.getvalue(), "invalid.xlsx")

    def test_full_aggregation_is_not_limited_by_detail_sample(self):
        orchestrator = Orchestrator(self.memory, self.root, max_filter_rows=3, max_retries=0)
        rows = [
            {"date": "2026-01-01", "product": "A", "channel": "线上", "amount": 100, "quantity": 1},
            {"date": "2026-01-02", "product": "A", "channel": "线上", "amount": 200, "quantity": 2},
            {"date": "2026-01-03", "product": "B", "channel": "门店", "amount": 300, "quantity": 3},
            {"date": "2026-02-01", "product": "B", "channel": "门店", "amount": 400, "quantity": 4},
            {"date": "2026-02-02", "product": "C", "channel": "线上", "amount": 500, "quantity": 5},
        ]
        result = orchestrator.analyze(project_name="全量聚合", request="分析季节趋势", data=rows)
        self.assertEqual(result["prepared_summary"]["rows"], 5)
        self.assertEqual(result["prepared_summary"]["sampled_rows"], 3)
        self.assertEqual(result["prepared_summary"]["total_amount"], 1500)
        self.assertEqual(result["charts"]["date_totals"], {"2026-01": 600.0, "2026-02": 900.0})
        self.assertEqual(result["insights"]["period_count"], 2)
        self.assertTrue(result["report_validation"]["checks"]["full_aggregate_total_matches"])

    def test_marketing_strategy_is_model_generated_and_structured(self):
        orchestrator = Orchestrator(self.memory, self.root, max_retries=0, llm=FixedLLM())
        rows = [
            {"date": "2026-01-01", "product": "A", "amount": 100},
            {"date": "2026-02-01", "product": "A", "amount": 180},
        ]
        result = orchestrator.analyze(project_name="模型策略", request="生成通俗的营销建议", data=rows)
        strategy = result["strategy"]
        self.assertTrue(strategy["llm_used"])
        self.assertEqual(strategy["generation_mode"], "model")
        self.assertEqual(strategy["actions"][0]["title"], "重点产品组合测试")
        markdown = Path(result["report"]["path"]).read_text(encoding="utf-8")
        self.assertIn("数据依据", markdown)
        self.assertIn("执行步骤", markdown)
        self.assertNotIn("模型补充建议", markdown)
        self.assertNotIn("通过 未分类", markdown)

    def test_marketing_fallback_does_not_invent_channel_or_audience(self):
        result = self.orchestrator.analyze(
            project_name="离线策略",
            request="生成营销建议",
            data=[{"date": "2026-01-01", "product": "A", "amount": 100}],
        )
        strategy = result["strategy"]
        self.assertFalse(strategy["llm_used"])
        self.assertEqual(strategy["generation_mode"], "offline_fallback")
        rendered = str(strategy)
        self.assertNotIn("高意向人群", rendered)
        self.assertNotIn("未转化访客", rendered)
        self.assertNotIn("未分类渠道", rendered)
        self.assertTrue(all(action.get("title") and action.get("why") and action.get("steps") and action.get("metrics") for action in strategy["actions"]))


if __name__ == "__main__": unittest.main()
