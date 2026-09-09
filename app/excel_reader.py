from __future__ import annotations

from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable, Sequence


EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xls"}
DATE_COLUMNS = {"date", "日期", "time", "时间", "invoicedate", "orderdate", "order_date", "订单日期"}
AMOUNT_COLUMNS = {"amount", "sales", "revenue", "销售额", "金额"}
QUANTITY_COLUMNS = {"quantity", "qty", "销量", "数量"}
PRICE_COLUMNS = {"price", "unitprice", "unit_price", "单价"}


def _plain_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def _headers(values: Sequence[Any]) -> list[str]:
    last = max((index for index, value in enumerate(values) if value not in (None, "")), default=-1)
    if last < 0:
        return []
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values[: last + 1], start=1):
        base = str(value).strip() if value not in (None, "") else f"列{index}"
        counts[base] = counts.get(base, 0) + 1
        headers.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return headers


def _records(header: list[str], rows: Iterable[Sequence[Any]]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for values in rows:
        selected = list(values[: len(header)])
        if not any(value not in (None, "") for value in selected):
            continue
        selected.extend([None] * (len(header) - len(selected)))
        records.append({name: _plain_value(value) for name, value in zip(header, selected)})
    return records


def _read_xlsx(content: bytes) -> tuple[list[dict[str, Any]], str]:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("无法读取 Excel 文件，请确认文件未损坏且格式正确") from exc
    try:
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            for values in iterator:
                header = _headers(values)
                if not header:
                    continue
                rows = _records(header, iterator)
                if rows:
                    return rows, sheet.title
                break
    finally:
        workbook.close()
    raise ValueError("Excel 中没有可分析的数据行")


def _xls_value(book: Any, cell: Any) -> Any:
    from xlrd import XL_CELL_DATE, xldate_as_datetime

    if cell.ctype == XL_CELL_DATE:
        return xldate_as_datetime(cell.value, book.datemode).isoformat()
    return cell.value


def _read_xls(content: bytes) -> tuple[list[dict[str, Any]], str]:
    try:
        import xlrd

        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    except Exception as exc:
        raise ValueError("无法读取旧版 .xls 文件，请确认文件未损坏且格式正确") from exc
    try:
        for sheet in workbook.sheets():
            header_row = -1
            header: list[str] = []
            for row_index in range(sheet.nrows):
                values = [_xls_value(workbook, cell) for cell in sheet.row(row_index)]
                header = _headers(values)
                if header:
                    header_row = row_index
                    break
            if header_row < 0:
                continue
            source_rows = (
                [_xls_value(workbook, cell) for cell in sheet.row(row_index)]
                for row_index in range(header_row + 1, sheet.nrows)
            )
            rows = _records(header, source_rows)
            if rows:
                return rows, sheet.name
    finally:
        workbook.release_resources()
    raise ValueError("Excel 中没有可分析的数据行")


def _validate_sales_columns(rows: list[dict[str, Any]]) -> None:
    lowered = {str(key).strip().lower() for key in rows[0]}
    missing: list[str] = []
    if not lowered.intersection(DATE_COLUMNS):
        missing.append("日期（date、日期、time 或 时间）")
    has_amount = bool(lowered.intersection(AMOUNT_COLUMNS))
    can_derive_amount = bool(lowered.intersection(QUANTITY_COLUMNS) and lowered.intersection(PRICE_COLUMNS))
    if not has_amount and not can_derive_amount:
        missing.append("销售额，或同时提供销量与单价列")
    if missing:
        raise ValueError("Excel 缺少必要列：" + "；".join(missing))


def read_excel_rows(content: bytes, filename: str) -> tuple[list[dict[str, Any]], str]:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in EXCEL_EXTENSIONS:
        raise ValueError("只支持 .xlsx、.xlsm 或 .xls 格式的 Excel 文件")
    if not content:
        raise ValueError("上传的 Excel 文件为空")
    rows, sheet_name = _read_xls(content) if suffix == ".xls" else _read_xlsx(content)
    _validate_sales_columns(rows)
    return rows, sheet_name
